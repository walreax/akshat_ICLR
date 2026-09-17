# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches

SRC = r"C:\Users\Akshat Jha\Downloads\write-up-longer.docx"
IMG_DIR = r"C:\Users\Akshat Jha\akshat_ICLR\writeups\example_images"

images = [f"{IMG_DIR}\\example_{i:02d}.jpg" for i in range(13)]

doc = docx.Document(SRC)
paras = doc.paragraphs

heading_idx = None
for i, p in enumerate(paras):
    if p.text.strip() == "3.6 Example of the Prompt and the Experimental Setup":
        heading_idx = i
        break

if heading_idx is None:
    raise RuntimeError("Could not find the 3.6 section heading; aborting without changes.")

# Collect the blank placeholder paragraphs right after the heading, up to the
# next non-blank paragraph (the "4. Results" heading).
blank_slots = []
j = heading_idx + 1
while j < len(paras) and paras[j].text.strip() == "":
    blank_slots.append(paras[j])
    j += 1

if len(blank_slots) < 5:
    raise RuntimeError(f"Only {len(blank_slots)} blank paragraph(s) available after the heading; "
                        "not enough room, aborting without changes.")

per_row = 3
img_idx = 0
slot_idx = 0
while img_idx < len(images):
    p = blank_slots[slot_idx]
    p.text = ""  # ensure clean
    row_imgs = images[img_idx: img_idx + per_row]
    for k, img_path in enumerate(row_imgs):
        run = p.add_run()
        run.add_picture(img_path, width=Inches(1.95))
        if k != len(row_imgs) - 1:
            p.add_run("  ")
    img_idx += per_row
    slot_idx += 1

doc.save(SRC)
print("saved, rows used:", slot_idx, "of", len(blank_slots), "available blank slots")
