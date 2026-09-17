# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt

SRC = r"C:\Users\Akshat Jha\Downloads\write-up-longer.docx"

# old caption text -> new caption text
RELABEL = {
    "seed_12": "Seed = 12",
    "seed_45": "Seed = 45",
    "seed_42": "Seed = 42",
    "seed_1000": "Seed = 1000",
    "seed45block20": "Block 20 @ Seed = 45",
    "seed45block16": "Block 16 @ Seed = 45",
    "seed45block12": "Block 12 @ Seed = 45",
    "seed45block8": "Block 8 @ Seed = 45",
    "seed45block4": "Block 4 @ Seed = 45",
    "block20seed42": "Block 20 @ Seed = 42",
    "block16seed42": "Block 16 @ Seed = 42",
    "block4seed42": "Block 4 @ Seed = 42",
    "block12seed42": "Block 12 @ Seed = 42",
    "block-8seed12": "Block 8 @ Seed = 12",
}

doc = docx.Document(SRC)
changed = 0
for table in doc.tables[:2]:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                text = p.text.strip()
                if text in RELABEL:
                    new_text = RELABEL[text]
                    # replace run text in place, keep italic/size formatting
                    for r in p.runs:
                        r.text = ""
                    if p.runs:
                        p.runs[0].text = new_text
                    else:
                        run = p.add_run(new_text)
                        run.italic = True
                        run.font.size = Pt(9)
                    changed += 1

doc.save(SRC)
print("captions relabeled:", changed)
