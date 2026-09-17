# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:\Users\Akshat Jha\Downloads\write-up-longer.docx"
IMG_DIR = r"C:\Users\Akshat Jha\akshat_ICLR\writeups\example_images"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

GENERATED = [
    ("example_09.jpg", "seed_12"),
    ("example_10.jpg", "seed_45"),
    ("example_11.jpg", "seed_42"),
    ("example_12.jpg", "seed_1000"),
]

ATTENTION = [
    ("example_00.jpg", "seed45block20"),
    ("example_01.jpg", "seed45block16"),
    ("example_02.jpg", "seed45block12"),
    ("example_03.jpg", "seed45block8"),
    ("example_04.jpg", "seed45block4"),
    ("example_05.jpg", "block20seed42"),
    ("example_06.jpg", "block16seed42"),
    ("example_07.jpg", "block4seed42"),
    ("example_08.jpg", "block12seed42"),
    ("block_8seed12.jpg", "block-8seed12"),
]


def round_picture_corners(run, adj="20000"):
    prst_geoms = run._r.findall(f".//{{{A_NS}}}prstGeom")
    for pg in prst_geoms:
        pg.set("prst", "roundRect")
        for child in list(pg):
            pg.remove(child)
        av_lst = OxmlElement("a:avLst")
        gd = OxmlElement("a:gd")
        gd.set("name", "adj")
        gd.set("fmla", f"val {adj}")
        av_lst.append(gd)
        pg.append(av_lst)


def build_grid(doc, anchor_el, items, cols):
    rows = (len(items) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    anchor_el.addnext(table._tbl)

    idx = 0
    for r in range(rows):
        for c in range(cols):
            if idx >= len(items):
                break
            fname, label = items[idx]
            cell = table.cell(r, c)
            cell.text = ""
            img_p = cell.paragraphs[0]
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = img_p.add_run()
            run.add_picture(f"{IMG_DIR}\\{fname}", width=Inches(1.7))
            round_picture_corners(run)

            cap_p = cell.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_before = Pt(2)
            cap_run = cap_p.add_run(label)
            cap_run.italic = True
            cap_run.font.size = Pt(9)
            idx += 1
    return table._tbl


doc = docx.Document(SRC)

# Remove the previously inserted image grid (the first table in the doc,
# which sits in section 3.6 before the real benchmark table in section 4.1).
old_table = doc.tables[0]
old_tbl_el = old_table._tbl
parent = old_tbl_el.getparent()
insertion_anchor = old_tbl_el.getprevious()  # the paragraph right before it
parent.remove(old_tbl_el)

# insertion_anchor is expected to be the (now empty) paragraph slot that
# used to precede the table -- use its element as the anchor to build after.
anchor_el = insertion_anchor

label1 = doc.add_paragraph()
label1.paragraph_format.space_before = Pt(6)
label1.paragraph_format.space_after = Pt(4)
r1 = label1.add_run("Generated Examples")
r1.bold = True
r1.font.size = Pt(11)
anchor_el.addnext(label1._p)
anchor_el = label1._p

tbl1_el = build_grid(doc, anchor_el, GENERATED, cols=4)
anchor_el = tbl1_el

label2 = doc.add_paragraph()
label2.paragraph_format.space_before = Pt(12)
label2.paragraph_format.space_after = Pt(4)
r2 = label2.add_run("Attention Maps")
r2.bold = True
r2.font.size = Pt(11)
anchor_el.addnext(label2._p)
anchor_el = label2._p

build_grid(doc, anchor_el, ATTENTION, cols=5)

doc.save(SRC)
print("saved.")
